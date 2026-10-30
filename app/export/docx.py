from pathlib import Path
from docx import Document
from app.export.common import sections


def save_docx(result, path: str | Path) -> None:
    get = (lambda key, default="": result.get(key, default)) if isinstance(result, dict) else (lambda key, default="": getattr(result, key, default))
    doc = Document()
    doc.add_heading(get("title", "PDF Summary"), level=0)
    p = doc.add_paragraph()
    p.add_run(f"Pages: {get('page_count', '')} | Language: {get('language', '')} | Processing time: {get('processing_seconds', '')} seconds")
    for title, body in sections(result):
        doc.add_heading(title, level=1)
        for line in (body or "Not provided.").splitlines():
            if line.strip().startswith(("- ", "• ")):
                doc.add_paragraph(line.strip()[2:], style="List Bullet")
            elif line.strip():
                doc.add_paragraph(line.strip())
    doc.save(Path(path))
